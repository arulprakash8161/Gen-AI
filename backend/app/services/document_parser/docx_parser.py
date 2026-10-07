import io
import uuid
from typing import Optional, Dict, Any, List
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from backend.app.models.document import ProcessedDocument, PageContent
from backend.app.services.document_parser.cleaner import clean_text
from backend.app.core.exceptions import AppException
from backend.app.core.logging import get_logger

logger = get_logger(__name__)


def parse_docx(
    file_bytes: bytes,
    filename: str = "document.docx",
    doc_id: Optional[str] = None,
) -> ProcessedDocument:
    """
    Extracts text from a DOCX byte stream, including body paragraphs and tables.
    """
    if not file_bytes:
        raise AppException(
            message=f"DOCX file '{filename}' is empty.",
            status_code=400,
            code="EMPTY_FILE",
        )

    document_id = doc_id or f"doc_{uuid.uuid4().hex[:12]}"

    try:
        doc = Document(io.BytesIO(file_bytes))
    except (PackageNotFoundError, Exception) as exc:
        logger.error(f"Failed to read DOCX '{filename}': {str(exc)}")
        raise AppException(
            message=f"Could not parse DOCX file '{filename}'. The file may be corrupt or invalid.",
            status_code=400,
            code="DOCX_PARSE_ERROR",
            details=str(exc),
        )

    extracted_blocks: List[str] = []

    # Extract paragraphs
    for para in doc.paragraphs:
        cleaned = clean_text(para.text)
        if cleaned:
            extracted_blocks.append(cleaned)

    # Extract tables (preserving tabular structure via pipes)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [clean_text(cell.text) for cell in row.cells]
            # Deduplicate contiguous repeated cells caused by merged cells
            deduped_cells = []
            for cell in row_cells:
                if not deduped_cells or cell != deduped_cells[-1]:
                    deduped_cells.append(cell)
            if any(deduped_cells):
                extracted_blocks.append(" | ".join(deduped_cells))

    if not extracted_blocks:
        logger.warning(f"DOCX '{filename}' contains no readable text content.")

    full_text = "\n\n".join(extracted_blocks)

    # Collect metadata from core properties if present
    docx_meta: Dict[str, Any] = {}
    try:
        core_props = doc.core_properties
        if core_props.title:
            docx_meta["title"] = core_props.title
        if core_props.author:
            docx_meta["author"] = core_props.author
        if core_props.subject:
            docx_meta["subject"] = core_props.subject
    except Exception:
        pass

    docx_meta["total_paragraphs"] = len(doc.paragraphs)
    docx_meta["total_tables"] = len(doc.tables)

    # DOCX files do not have explicit page markers in raw XML; represent content with 1-indexed section/page
    pages = [
        PageContent(
            page_number=1,
            text=full_text,
            character_count=len(full_text),
        )
    ]

    return ProcessedDocument(
        doc_id=document_id,
        filename=filename,
        file_type="docx",
        total_pages=1,
        text=full_text,
        character_count=len(full_text),
        pages=pages,
        metadata=docx_meta,
    )
