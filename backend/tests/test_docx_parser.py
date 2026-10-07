import io
import pytest
from docx import Document
from backend.app.services.document_parser.docx_parser import parse_docx
from backend.app.core.exceptions import AppException


def create_sample_docx_bytes(paragraphs: list[str], table_data: list[list[str]] = None) -> bytes:
    """Helper to generate a real DOCX file in memory."""
    doc = Document()
    doc.core_properties.title = "Test DOCX Document"
    doc.core_properties.author = "SIH Engineering"

    for p in paragraphs:
        doc.add_paragraph(p)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                table.cell(r_idx, c_idx).text = cell_value

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_docx_valid():
    docx_bytes = create_sample_docx_bytes(
        paragraphs=["Executive overview paragraph.", "Second key finding."],
        table_data=[["Metric", "Value"], ["Accuracy", "99%"]]
    )

    doc = parse_docx(file_bytes=docx_bytes, filename="overview.docx")

    assert doc.filename == "overview.docx"
    assert doc.file_type == "docx"
    assert doc.total_pages == 1
    assert "Executive overview paragraph." in doc.text
    assert "Second key finding." in doc.text
    assert "Metric | Value" in doc.text
    assert "Accuracy | 99%" in doc.text
    assert doc.metadata.get("title") == "Test DOCX Document"
    assert doc.metadata.get("author") == "SIH Engineering"


def test_parse_docx_empty_bytes_raises_error():
    with pytest.raises(AppException) as exc_info:
        parse_docx(file_bytes=b"", filename="empty.docx")
    assert exc_info.value.code == "EMPTY_FILE"
    assert exc_info.value.status_code == 400


def test_parse_docx_invalid_bytes_raises_error():
    with pytest.raises(AppException) as exc_info:
        parse_docx(file_bytes=b"Invalid non-docx content", filename="corrupt.docx")
    assert exc_info.value.code == "DOCX_PARSE_ERROR"
    assert exc_info.value.status_code == 400
